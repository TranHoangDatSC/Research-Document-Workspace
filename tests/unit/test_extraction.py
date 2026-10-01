"""Text extraction: the pure .txt/.pdf/.docx extractors, then the
"extract text" action on a stored document (file -> text -> MongoDB).
"""
import io
import unittest
import zipfile
from unittest.mock import patch
from uuid import UUID

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter

from support import FakeBackend

from app.services import documents as documents_service
from app.extractors import MAX_EXTRACTED_CHARACTERS, MAX_INPUT_BYTES, ExtractionError, extract_text


def build_pdf(text="Hello PDF extraction"):
    """Hand-built single-page PDF (no external renderer) with a real text stream."""
    content = f"BT /F1 24 Tf 20 100 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> "
        b"/MediaBox [0 0 200 200] /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n".encode("latin-1") + content + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode("latin-1") + body + b"\nendobj\n"
    xref_offset = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode("latin-1")
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode("latin-1")
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF"
    ).encode("latin-1")
    return bytes(out)


def build_blank_pdf(password=None):
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    if password:
        writer.encrypt(password)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def build_docx(paragraphs=("Hello DOCX extraction",)):
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class ExtractorTestCase(unittest.TestCase):
    def assert_extraction_error(self, code, content, extension):
        with self.assertRaises(ExtractionError) as caught:
            extract_text(content, extension)
        self.assertEqual(caught.exception.code, code)


class PlainTextExtractionTests(ExtractorTestCase):
    def test_utf8_vietnamese(self):
        text = "Tài liệu nghiên cứu\nCloud Docker"
        result = extract_text(text.encode("utf-8"), ".txt")
        self.assertEqual(result.text, text)
        self.assertEqual(result.method, "plain_text")
        self.assertEqual((result.character_count, result.word_count, result.truncated), (len(text), 6, False))

    def test_bom_and_line_endings_normalized(self):
        result = extract_text(b"\xef\xbb\xbfFirst\r\nSecond\rThird", " .TXT ")
        self.assertEqual(result.text, "First\nSecond\nThird")
        self.assertEqual(result.word_count, 3)

    def test_empty_and_whitespace_give_empty_text(self):
        for content in (b"", b" \t\r\n", b"\xef\xbb\xbf"):
            with self.subTest(content=content):
                result = extract_text(content, ".txt")
                self.assertEqual((result.text, result.character_count, result.word_count, result.truncated), ("", 0, 0, False))

    def test_rejected_content(self):
        self.assert_extraction_error("invalid_utf8", b"\xff\xfe", ".txt")
        self.assert_extraction_error("invalid_text", b"hello\x00world", ".txt")
        self.assert_extraction_error("input_too_large", b"a" * (MAX_INPUT_BYTES + 1), ".txt")
        self.assert_extraction_error("unsupported_format", b"hello", ".exe")

    def test_truncation_counts_saved_text(self):
        result = extract_text(b"one two three", ".txt", max_characters=7)
        self.assertEqual((result.text, result.character_count, result.word_count, result.truncated), ("one two", 7, 2, True))

    def test_exact_limit_is_not_truncated(self):
        result = extract_text(b"abcde", ".txt", max_characters=5)
        self.assertEqual((result.text, result.truncated), ("abcde", False))

    def test_default_limit(self):
        result = extract_text(b"a" * (MAX_EXTRACTED_CHARACTERS + 1), ".txt")
        self.assertEqual(result.character_count, MAX_EXTRACTED_CHARACTERS)
        self.assertTrue(result.truncated)

    def test_invalid_arguments(self):
        for limit in (0, -1, True, 1.5, 200_001):
            with self.subTest(limit=limit), self.assertRaises(ValueError):
                extract_text(b"hello", ".txt", limit)
        with self.assertRaises(TypeError):
            extract_text("hello", ".txt")


class PdfExtractionTests(ExtractorTestCase):
    def test_pdf_text(self):
        result = extract_text(build_pdf("Hello PDF extraction"), ".pdf")
        self.assertEqual((result.text, result.method, result.truncated), ("Hello PDF extraction", "pdf_text", False))

    def test_uppercase_extension(self):
        self.assertEqual(extract_text(build_pdf("case"), " .PDF ").text, "case")

    def test_blank_page_is_empty_not_error(self):
        result = extract_text(build_blank_pdf(), ".pdf")
        self.assertEqual((result.text, result.character_count), ("", 0))

    def test_rejected_pdfs(self):
        self.assert_extraction_error("encrypted_file", build_blank_pdf(password="secret"), ".pdf")
        self.assert_extraction_error("corrupt_file", b"not a pdf file", ".pdf")
        self.assert_extraction_error("input_too_large", b"a" * (MAX_INPUT_BYTES + 1), ".pdf")

    def test_truncation(self):
        result = extract_text(build_pdf("one two three"), ".pdf", max_characters=7)
        self.assertEqual((result.text, result.truncated), ("one two", True))


class DocxExtractionTests(ExtractorTestCase):
    def test_docx_text(self):
        result = extract_text(build_docx(["First paragraph", "Second paragraph"]), ".docx")
        self.assertEqual((result.text, result.method, result.truncated), ("First paragraph\nSecond paragraph", "docx_text", False))

    def test_uppercase_extension(self):
        self.assertEqual(extract_text(build_docx(["case"]), " .DOCX ").text, "case")

    def test_no_paragraphs_is_empty_not_error(self):
        result = extract_text(build_docx([]), ".docx")
        self.assertEqual((result.text, result.character_count), ("", 0))

    def test_rejected_docx(self):
        self.assert_extraction_error("corrupt_file", b"not a docx file", ".docx")
        # A well-formed ZIP that is not a Word package must still be rejected cleanly.
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("hello.txt", "not a docx package")
        self.assert_extraction_error("corrupt_file", buffer.getvalue(), ".docx")
        self.assert_extraction_error("input_too_large", b"a" * (MAX_INPUT_BYTES + 1), ".docx")

    def test_truncation(self):
        result = extract_text(build_docx(["one two three"]), ".docx", max_characters=7)
        self.assertEqual((result.text, result.truncated), ("one two", True))


def build_pptx():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])  # title + content
    slide.shapes.title.text = "Workload forecasting"
    slide.placeholders[1].text = "ARIMA vs LSTM"
    slide.notes_slide.notes_text_frame.text = "Mention the 2021 survey"
    second = presentation.slides.add_slide(presentation.slide_layouts[5])  # title only
    second.shapes.title.text = "Results"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Model", "MAPE"
    table.cell(1, 0).text, table.cell(1, 1).text = "LSTM", "7.2%"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def build_zip(files):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return buffer.getvalue()


class ZipExtractionTests(ExtractorTestCase):
    def test_lists_everything_and_reads_what_it_can(self):
        content = build_zip({
            "notes/readme.txt": b"Plain text inside",
            "slides/deck.pptx": build_pptx(),
            "report.docx": build_docx(["Docx inside"]),
            "photo.png": b"\x89PNG not read locally",
            "inner.zip": build_zip({"deep.txt": b"never opened"}),
            "broken.pdf": b"not a pdf",
        })
        result = extract_text(content, ".zip")
        self.assertEqual(result.method, "zip_text")
        self.assertIn("--- Nội dung tệp nén (6 tệp) ---", result.text)
        for name in ("notes/readme.txt", "photo.png", "inner.zip"):
            self.assertIn(f"- {name} (", result.text)  # listed
        self.assertIn("Plain text inside", result.text)
        self.assertIn("ARIMA vs LSTM", result.text)
        self.assertIn("Docx inside", result.text)
        self.assertNotIn("never opened", result.text)       # nested zips are not opened
        self.assertIn("--- broken.pdf ---\n(không đọc được)", result.text)

    def test_decompression_budget_stops_a_zip_bomb(self):
        from app import extractors
        content = build_zip({"a.txt": b"A" * 1000, "b.txt": b"B" * 1000})
        with patch.object(extractors, "ZIP_MAX_UNCOMPRESSED", 1500):
            text = extract_text(content, ".zip").text
        self.assertIn("A" * 1000, text)
        self.assertIn("--- b.txt ---\n(bỏ qua: quá lớn để đọc trong tệp nén)", text)

    def test_corrupt_zip(self):
        self.assert_extraction_error("corrupt_file", b"not a zip", ".zip")


class PresentationAndDataExtractionTests(ExtractorTestCase):
    def test_pptx_slides_tables_and_notes(self):
        result = extract_text(build_pptx(), ".pptx")
        self.assertEqual(result.method, "pptx_text")
        for fragment in ("--- Slide 1 ---", "Workload forecasting", "ARIMA vs LSTM",
                         "Ghi chú: Mention the 2021 survey", "--- Slide 2 ---", "Model\tMAPE", "LSTM\t7.2%"):
            self.assertIn(fragment, result.text)
        self.assertLess(result.text.index("Slide 1"), result.text.index("Slide 2"))

    def test_corrupt_pptx(self):
        self.assert_extraction_error("corrupt_file", b"not a pptx", ".pptx")

    def test_xlsx_sheets_rows_and_computed_values(self):
        workbook = Workbook()
        first = workbook.active
        first.title = "Kết quả"
        first.append(["Model", "MAPE", None])
        first.append(["LSTM", 7.2])
        first.append([])                 # blank row skipped
        first.append(["Tổng", "=1+1"])   # formula: no cached value when written by openpyxl
        workbook.create_sheet("Trống")    # empty sheet skipped
        workbook.create_sheet("Ghi chú").append(["Nguồn: khảo sát 2021"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        result = extract_text(buffer.getvalue(), ".xlsx")
        self.assertEqual(result.method, "xlsx_text")
        self.assertEqual(
            result.text,
            "--- Sheet: Kết quả ---\nModel\tMAPE\nLSTM\t7.2\nTổng\n\n--- Sheet: Ghi chú ---\nNguồn: khảo sát 2021",
        )

    def test_corrupt_xlsx(self):
        self.assert_extraction_error("corrupt_file", b"not a workbook", ".xlsx")

    def test_markdown_csv_json_read_as_text(self):
        for extension, content in ((".md", "# Tiêu đề\n- ý một"), (".csv", "model,mape\nlstm,7.2"), (".json", '{"model": "lstm"}')):
            with self.subTest(extension=extension):
                result = extract_text(content.encode("utf-8"), extension)
                self.assertEqual((result.text, result.method), (content, "plain_text"))


class ExtractDocumentTests(unittest.TestCase):
    """The "Trích xuất văn bản" action on an uploaded document."""

    def setUp(self):
        self.backend = FakeBackend().install(self)
        self.client = self.backend.client(self)
        self.project_id = self.client.post("/projects", json={"name": "P"}).json()["id"]

    def upload(self, name, content):
        response = self.client.post(f"/projects/{self.project_id}/documents", files={"file": (name, content)})
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()["id"]

    def extract(self, document_id):
        return self.client.post(f"/documents/{document_id}/extract")

    def test_txt_pdf_docx_text_saved_to_metadata(self):
        for name, content, text, method in (
            ("a.txt", "Xin chào Cloud Docker".encode("utf-8"), "Xin chào Cloud Docker", "plain_text"),
            ("b.pdf", build_pdf("From a PDF"), "From a PDF", "pdf_text"),
            ("c.docx", build_docx(["From a DOCX"]), "From a DOCX", "docx_text"),
        ):
            with self.subTest(name=name):
                document_id = self.upload(name, content)
                response = self.extract(document_id)
                self.assertEqual(response.status_code, 200)
                extracted = response.json()["extracted_text"]
                self.assertEqual((extracted["text"], extracted["method"]), (text, method))
                self.assertIn("extracted_at", extracted)
                # What the API returns is exactly what was stored.
                self.assertEqual(self.backend.details[document_id]["extracted_text"], extracted)

    def test_page_button_extracts_and_shows_text(self):
        document_id = self.upload("a.txt", b"Visible extracted sentence")
        response = self.client.post(f"/ui/documents/{document_id}/extract", follow_redirects=False)
        # Back on the extracted-text tab, not the default info tab.
        self.assertEqual(response.headers["location"], f"/ui/documents/{document_id}#extract")
        self.assertIn("Visible extracted sentence", self.client.get(response.headers["location"]).text)

    def test_reextraction_overwrites(self):
        document_id = self.upload("a.txt", b"first version")
        self.extract(document_id)
        stored = self.backend.details[document_id]
        self.backend.objects[f"documents/{document_id}/original.txt"] = b"second version"
        self.backend.documents[UUID(document_id)]["size_bytes"] = len(b"second version")
        self.extract(document_id)
        self.assertEqual(stored["extracted_text"]["text"], "second version")

    def test_unreadable_file_is_422_and_saves_nothing(self):
        for name, content in (("bad.pdf", b"not a pdf"), ("locked.pdf", build_blank_pdf(password="secret"))):
            with self.subTest(name=name):
                document_id = self.upload(name, content)
                self.assertEqual(self.extract(document_id).status_code, 422)
                self.assertIsNone(self.backend.details[document_id]["extracted_text"])

    def test_pptx_document(self):
        document_id = self.upload("deck.pptx", build_pptx())
        extracted = self.extract(document_id).json()["extracted_text"]
        self.assertIn("ARIMA vs LSTM", extracted["text"])

    def test_text_files_are_extracted_on_upload(self):
        document_id = self.upload("notes.md", "# Ghi chú\nARIMA".encode())
        self.assertEqual(self.backend.details[document_id]["extracted_text"]["text"], "# Ghi chú\nARIMA")

    def test_zip_document(self):
        document_id = self.upload("bundle.zip", build_zip({"readme.txt": b"Inside the zip"}))
        extracted = self.backend.details[document_id]["extracted_text"]
        self.assertEqual(extracted["method"], "zip_text")
        self.assertIn("Inside the zip", extracted["text"])

    def test_too_large_to_extract_is_refused_before_download(self):
        document_id = self.upload("big.txt", b"x" * 20)
        self.backend.downloads.clear()  # the upload's own extraction read it once
        with patch.object(documents_service, "EXTRACT_MAX_BYTES", 10):
            response = self.extract(document_id)
        self.assertEqual(response.status_code, 422)
        self.assertEqual(self.backend.downloads, [])

    def test_media_page_offers_ai_analysis(self):
        document_id = self.upload("clip.mp4", b"x")
        html = self.client.get(f"/ui/documents/{document_id}").text
        self.assertIn("data-inline-extract", html)
        self.assertIn("Phân tích bằng AI", html)
        self.assertIn("tệp rời khỏi máy chủ của bạn", html)  # the privacy trade-off is stated

    def test_not_ready_document_is_409(self):
        document_id = self.upload("a.txt", b"hello")
        self.backend.documents[UUID(document_id)]["status"] = "pending"
        self.assertEqual(self.extract(document_id).status_code, 409)

    def test_storage_failures_are_503(self):
        document_id = self.upload("a.txt", b"hello")
        for store in ("minio", "mongo"):
            with self.subTest(store=store):
                self.backend.fail = {store}
                self.assertEqual(self.extract(document_id).status_code, 503)
        self.backend.fail = set()
        self.assertEqual(self.backend.details[document_id]["extracted_text"]["text"], "hello")  # from upload, untouched


if __name__ == "__main__":
    unittest.main()
